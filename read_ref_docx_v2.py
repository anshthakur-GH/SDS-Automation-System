import docx
import io
import sys

# Ensure UTF-8 output
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def read_docx(path):
    doc = docx.Document(path)
    lines = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(f"P: {para.text}")
    for table in doc.tables:
        lines.append("--- TABLE ---")
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(row_data))
    return '\n'.join(lines)

if __name__ == "__main__":
    path = r'./'
    print(read_docx(path))
